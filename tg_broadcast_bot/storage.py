"""SQLite-хранилище контактов и импорт из Excel."""
import io
import re
import sqlite3
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook

NAME_HEADERS = ("фио", "имя", "клиент", "name", "full name", "fio")
PHONE_HEADERS = ("телефон", "тел", "номер", "phone", "mobile", "моб")


def normalize_phone(raw) -> str | None:
    """Приводит номер к виду 79991234567. Возвращает None, если номер не похож на телефон."""
    if raw is None:
        return None
    if isinstance(raw, float):
        raw = int(raw)
    digits = re.sub(r"\D", "", str(raw))
    if len(digits) == 11 and digits[0] == "8":
        digits = "7" + digits[1:]
    elif len(digits) == 10 and digits[0] == "9":
        digits = "7" + digits
    if not 10 <= len(digits) <= 15:
        return None
    return digits


def first_name(full_name: str) -> str:
    """Из «Иванов Иван Иванович» берёт «Иван»; из одного слова — само слово."""
    parts = full_name.split()
    if len(parts) >= 2:
        return parts[1]
    return parts[0] if parts else ""


@dataclass
class ImportResult:
    added: int
    updated: int
    skipped: list[str]


def _find_columns(header: list) -> tuple[int, int] | None:
    name_col = phone_col = None
    for i, cell in enumerate(header):
        text = str(cell or "").strip().lower()
        if name_col is None and any(text.startswith(h) for h in NAME_HEADERS):
            name_col = i
        elif phone_col is None and any(text.startswith(h) for h in PHONE_HEADERS):
            phone_col = i
    if name_col is None or phone_col is None:
        return None
    return name_col, phone_col


def parse_excel(data: bytes) -> tuple[list[tuple[str, str]], list[str]]:
    """Читает первый лист. Ищет колонки ФИО и Телефон по заголовку,
    иначе считает, что ФИО — первая колонка, телефон — вторая."""
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    cols = _find_columns(list(rows[0]))
    first_row = 1
    if cols:
        rows = rows[1:]
        first_row = 2
    else:
        cols = (0, 1)
    name_col, phone_col = cols

    contacts, skipped = [], []
    for n, row in enumerate(rows, start=first_row):
        row = list(row) + [None] * (max(cols) + 1 - len(row))
        name = str(row[name_col] or "").strip()
        phone = normalize_phone(row[phone_col])
        if not name and row[phone_col] is None:
            continue
        if not name or not phone:
            skipped.append(f"строка {n}: {name or '—'} / {row[phone_col] or '—'}")
            continue
        contacts.append((name, phone))
    return contacts, skipped


class Storage:
    def __init__(self, path: str):
        self.db = sqlite3.connect(path)
        self.db.row_factory = sqlite3.Row
        self.db.execute(
            """CREATE TABLE IF NOT EXISTS contacts (
                phone TEXT PRIMARY KEY,
                full_name TEXT NOT NULL,
                tg_id INTEGER,
                subscribed INTEGER NOT NULL DEFAULT 0
            )"""
        )
        self.db.commit()

    def import_contacts(self, contacts: list[tuple[str, str]], skipped: list[str]) -> ImportResult:
        added = updated = 0
        for name, phone in contacts:
            cur = self.db.execute(
                "UPDATE contacts SET full_name = ? WHERE phone = ?", (name, phone)
            )
            if cur.rowcount:
                updated += 1
            else:
                self.db.execute(
                    "INSERT INTO contacts (phone, full_name) VALUES (?, ?)", (phone, name)
                )
                added += 1
        self.db.commit()
        return ImportResult(added, updated, skipped)

    def link(self, phone: str, tg_id: int) -> str | None:
        """Привязывает Telegram-аккаунт к контакту из базы. Возвращает ФИО или None."""
        row = self.db.execute("SELECT full_name FROM contacts WHERE phone = ?", (phone,)).fetchone()
        if not row:
            return None
        self.db.execute("UPDATE contacts SET tg_id = NULL WHERE tg_id = ?", (tg_id,))
        self.db.execute(
            "UPDATE contacts SET tg_id = ?, subscribed = 1 WHERE phone = ?", (tg_id, phone)
        )
        self.db.commit()
        return row["full_name"]

    def set_subscribed(self, tg_id: int, value: bool) -> bool:
        cur = self.db.execute(
            "UPDATE contacts SET subscribed = ? WHERE tg_id = ?", (int(value), tg_id)
        )
        self.db.commit()
        return cur.rowcount > 0

    def recipients(self) -> list[sqlite3.Row]:
        return self.db.execute(
            "SELECT full_name, tg_id FROM contacts WHERE tg_id IS NOT NULL AND subscribed = 1"
        ).fetchall()

    def stats(self) -> dict:
        q = lambda sql: self.db.execute(sql).fetchone()[0]
        return {
            "total": q("SELECT COUNT(*) FROM contacts"),
            "linked": q("SELECT COUNT(*) FROM contacts WHERE tg_id IS NOT NULL"),
            "subscribed": q("SELECT COUNT(*) FROM contacts WHERE tg_id IS NOT NULL AND subscribed = 1"),
        }

    def export_unlinked(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Не подключены"
        ws.append(["ФИО", "Телефон"])
        for row in self.db.execute(
            "SELECT full_name, phone FROM contacts WHERE tg_id IS NULL ORDER BY full_name"
        ):
            ws.append([row["full_name"], "+" + row["phone"]])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def clear(self) -> None:
        self.db.execute("DELETE FROM contacts")
        self.db.commit()
